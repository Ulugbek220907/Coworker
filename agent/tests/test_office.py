"""Office module: atomic workbook writes, unique backups, scratch-only output, bounded COM.

Fixtures are generated with openpyxl and PyMuPDF, so no Office install is needed.
COM and LibreOffice are always replaced by fakes; no test starts a real Office.
"""
from __future__ import annotations

import os
import subprocess
import tempfile
import threading
import time
from pathlib import Path

import openpyxl
import pymupdf
import pytest

import coworker.config
from coworker import office


@pytest.fixture(autouse=True)
def home(tmp_path, monkeypatch) -> Path:
    """Each test gets its own COWORKER_HOME; no real AppData folder is touched."""
    root = tmp_path / "home"
    root.mkdir()
    monkeypatch.setenv("COWORKER_HOME", str(root))
    monkeypatch.setattr(coworker.config, "config_dir", lambda: root)
    return root


@pytest.fixture
def book(tmp_path) -> Path:
    """'Budget' with a formula, plus a Cyrillic-titled second sheet."""
    path = tmp_path / "docs" / "Budget.xlsx"
    path.parent.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"], ws["B1"] = "Item", "Amount"
    ws["A2"], ws["B2"] = "Rent", 1200
    ws["D2"] = "=B2*2"
    wb.create_sheet("Бюджет")["A1"] = "Qator"
    wb.save(path)
    return path


def make_pdf(path: Path, pages: int) -> Path:
    doc = pymupdf.open()
    for i in range(pages):
        doc.new_page().insert_text((72, 72), f"page {i + 1}")
    doc.save(path)
    doc.close()
    return path


def listing(folder: Path) -> set[str]:
    return {p.name for p in folder.iterdir()}


# ------------------------------------------------------- atomic workbook write

def test_failed_swap_leaves_original_intact(book, monkeypatch):
    original = book.read_bytes()
    before = listing(book.parent)

    def refuse(src, dst):
        raise OSError("file is open in Excel")

    monkeypatch.setattr(office.os, "replace", refuse)
    result = office.sheet_write(str(book), {"B2": 1500}, "Budget")

    assert result["code"] == "office_failed"
    assert "asl fayl o'zgarmadi" in result["error"]
    assert book.read_bytes() == original
    assert Path(result["backup"]).read_bytes() == original
    assert listing(book.parent) - before == {Path(result["backup"]).name}


def test_failed_flush_to_disk_leaves_original_intact(book, monkeypatch):
    original = book.read_bytes()
    real_fsync = os.fsync
    calls: list[int] = []

    def fsync(fd):
        calls.append(fd)
        if len(calls) == 2:          # call 1 is the backup copy, call 2 the replacement
            raise OSError("disk full")
        return real_fsync(fd)

    monkeypatch.setattr(office.os, "fsync", fsync)
    result = office.sheet_write(str(book), {"B2": 1500}, "Budget")

    assert "asl fayl o'zgarmadi" in result["error"]     # failed at the swap, not the backup
    assert Path(result["backup"]).exists()
    assert book.read_bytes() == original
    assert not [p for p in book.parent.iterdir() if p.suffix == ".tmp"]


def test_failed_save_leaves_original_intact(book, monkeypatch):
    original = book.read_bytes()

    def boom(self, *args, **kwargs):
        raise OSError("serialisation failed")

    monkeypatch.setattr(openpyxl.Workbook, "save", boom)
    result = office.sheet_write(str(book), {"B2": 1500}, "Budget")

    assert result["code"] == "office_failed"
    assert book.read_bytes() == original


def test_temp_file_is_created_beside_the_original(book, monkeypatch):
    seen: list[Path] = []
    real_mkstemp = tempfile.mkstemp

    def spy(**kwargs):
        seen.append(Path(kwargs["dir"]))
        return real_mkstemp(**kwargs)

    monkeypatch.setattr(office.tempfile, "mkstemp", spy)
    assert office.sheet_write(str(book), {"B2": 1}, "Budget")["ok"]
    assert seen == [book.parent]


def test_successful_write_changes_only_named_cells(book):
    original = book.read_bytes()
    result = office.sheet_write(str(book), {"B2": "1 500", "c3": "hello"}, "budget")

    assert result["ok"] is True
    assert result["sheet"] == "Budget"
    assert result["applied"] == {"B2": "1 500", "C3": "hello"}
    wb = openpyxl.load_workbook(book)
    assert wb["Budget"]["B2"].value == 1500
    assert wb["Budget"]["A2"].value == "Rent"
    assert wb["Budget"]["D2"].value == "=B2*2"          # formulas survive the rewrite
    assert wb["Бюджет"]["A1"].value == "Qator"
    assert wb["Budget"]["C3"].value == "hello"
    assert Path(result["backup"]).read_bytes() == original
    assert not [p for p in book.parent.iterdir() if p.suffix == ".tmp"]


def test_write_to_a_missing_sheet_changes_nothing(book):
    original = book.read_bytes()
    result = office.sheet_write(str(book), {"B2": 1}, "Nope")
    assert result["code"] == "arg_invalid"
    assert book.read_bytes() == original
    assert not any("backup" in p.name for p in book.parent.iterdir())


# --------------------------------------------------------------- backups

def test_backups_in_the_same_second_never_collide(book, monkeypatch):
    monkeypatch.setattr(office, "_stamp", lambda: "20261008-120000")
    first = office.sheet_write(str(book), {"B2": 1}, "Budget")
    second = office.sheet_write(str(book), {"B2": 2}, "Budget")

    assert first["ok"] and second["ok"]
    assert Path(first["backup"]).name == "Budget.backup-20261008-120000-1.xlsx"
    assert Path(second["backup"]).name == "Budget.backup-20261008-120000-2.xlsx"
    assert Path(first["backup"]).exists() and Path(second["backup"]).exists()


def test_an_existing_backup_name_is_never_overwritten(book, monkeypatch):
    monkeypatch.setattr(office, "_stamp", lambda: "20261008-120000")
    squatter = book.with_name("Budget.backup-20261008-120000-1.xlsx")
    squatter.write_bytes(b"keep me")

    result = office.sheet_write(str(book), {"B2": 9}, "Budget")

    assert squatter.read_bytes() == b"keep me"
    assert Path(result["backup"]).name == "Budget.backup-20261008-120000-2.xlsx"


# ------------------------------------------------------------ format refusals

@pytest.mark.parametrize("name", ["Book.xlsm", "Book.xlsb", "Book.xltm", "Book.xls", "Doc.docm"])
def test_macro_enabled_formats_are_refused_before_any_backup(tmp_path, name):
    target = tmp_path / name
    target.write_bytes(b"PK-not-really-a-workbook")
    before = listing(tmp_path)

    result = office.sheet_write(str(target), {"A1": 1})

    assert result["code"] == "arg_invalid"
    assert "Makrosli" in result["error"]
    assert target.read_bytes() == b"PK-not-really-a-workbook"
    assert listing(tmp_path) == before


@pytest.mark.parametrize("name", ["Template.xltx", "Data.csv"])
def test_other_non_xlsx_formats_are_refused(tmp_path, name):
    target = tmp_path / name
    target.write_bytes(b"x")
    result = office.sheet_write(str(target), {"A1": 1})
    assert result["code"] == "arg_invalid"
    assert "Faqat .xlsx" in result["error"]


# ------------------------------------------------------------------ reads

def test_sheet_list_names_every_sheet(book):
    result = office.sheet_list(str(book))
    assert [s["name"] for s in result["sheets"]] == ["Budget", "Бюджет"]
    assert result["sheets"][0]["rows"] == 2


def test_sheet_read_finds_a_cyrillic_sheet_by_a_loose_name(book):
    result = office.sheet_read(str(book), "бюджет")
    assert result["sheet"] == "Бюджет"
    assert "Qator" in result["table"]


def test_sheet_read_reports_truncation(tmp_path):
    path = tmp_path / "long.xlsx"
    wb = openpyxl.Workbook()
    for i in range(1, 61):
        wb.active.cell(row=i, column=1, value=i)
    wb.save(path)

    result = office.sheet_read(str(path))
    assert result["shown_rows"] == office.MAX_ROWS
    assert result["truncated"] is True


def test_sheet_read_unknown_sheet_lists_the_real_ones(book):
    result = office.sheet_read(str(book), "Nope")
    assert result["code"] == "arg_invalid"
    assert "Budget" in result["error"]


def test_reads_refuse_non_excel_files(tmp_path):
    notes = tmp_path / "notes.txt"
    notes.write_text("hello")
    assert office.sheet_list(str(notes))["code"] == "arg_invalid"
    assert office.sheet_read(str(notes))["code"] == "arg_invalid"


# ------------------------------------------------------------- COM timeouts

def test_overrunning_com_call_is_abandoned_and_counted():
    release = threading.Event()
    before = office.com_abandoned()

    def stuck():
        release.wait(10)
        return {"ok": True}

    started = time.monotonic()
    try:
        result = office._run_bounded(stuck, 0.2)
    finally:
        release.set()

    assert time.monotonic() - started < 5
    assert result["code"] == "timeout"
    assert office.com_abandoned() == before + 1


def test_com_call_within_the_limit_returns_its_result():
    assert office._run_bounded(lambda: {"ok": True, "path": "x"}, 5) == {"ok": True, "path": "x"}


def test_com_exception_becomes_an_error_not_a_crash():
    def broken():
        raise RuntimeError("rpc server unavailable")

    result = office._run_bounded(broken, 5)
    assert result["code"] == "office_failed"
    assert "rpc server" in result["error"]


def test_to_pdf_does_not_fall_back_to_libreoffice_after_an_office_overrun(tmp_path, monkeypatch):
    src = tmp_path / "letter.docx"
    src.write_bytes(b"docx")
    release = threading.Event()

    def stuck(source, target):
        release.wait(10)
        return {"ok": True}

    def no_fallback(*_args):
        raise AssertionError("LibreOffice must not run after an Office overrun")

    monkeypatch.setattr(office, "com_apps", lambda: ["Word"])
    monkeypatch.setattr(office, "COM_TIMEOUT", 0.2)
    monkeypatch.setattr(office, "_pdf_via_com", stuck)
    monkeypatch.setattr(office, "_pdf_via_libreoffice", no_fallback)
    try:
        result = office.to_pdf(str(src))
    finally:
        release.set()

    assert result["code"] == "timeout"
    assert "path" not in result


# ----------------------------------------------------------- LibreOffice

class FakePopen:
    """Stands in for subprocess.Popen: the child never finishes on its own."""

    instances: list["FakePopen"] = []

    def __init__(self, argv, **kwargs):
        self.argv = argv
        self.kwargs = kwargs
        self.pid = 4242
        self.returncode = None
        self.stderr = None
        FakePopen.instances.append(self)

    def communicate(self, timeout=None):
        raise subprocess.TimeoutExpired(self.argv, timeout)

    def wait(self, timeout=None):
        return 0

    def kill(self):
        pass


def test_libreoffice_overrun_kills_the_tree_and_cleans_its_folder(tmp_path, home, monkeypatch):
    src = tmp_path / "memo.rtf"
    src.write_bytes(b"{\\rtf1 x}")
    killed: list[int] = []
    FakePopen.instances.clear()
    monkeypatch.setenv("COWORKER_TEST_API_KEY", "must-not-leak")
    monkeypatch.setenv("COWORKER_PLAIN_VALUE", "kept")
    monkeypatch.setattr(office, "com_apps", lambda: [])
    monkeypatch.setattr(office, "libreoffice_path", lambda: "C:/fake/soffice.exe")
    monkeypatch.setattr(office, "CONVERT_TIMEOUT", 0.1)
    monkeypatch.setattr(office, "_kill_tree", lambda proc: killed.append(proc.pid))
    monkeypatch.setattr(office.subprocess, "Popen", FakePopen)

    result = office.to_pdf(str(src))

    assert result["code"] == "timeout"
    assert killed == [4242]
    env = FakePopen.instances[0].kwargs["env"]
    assert "COWORKER_TEST_API_KEY" not in env
    assert env["COWORKER_PLAIN_VALUE"] == "kept"
    scratch = home / "scratch"
    assert not [p for p in scratch.rglob("*") if p.name.startswith(".lo-")]


def test_to_pdf_uses_libreoffice_when_office_is_missing(tmp_path, home, monkeypatch):
    src = tmp_path / "letter.docx"
    src.write_bytes(b"docx")

    def fake_child(argv, *, cwd, timeout):
        out_dir = Path(argv[argv.index("--outdir") + 1])
        (out_dir / "letter.pdf").write_bytes(b"%PDF-1.4 fake")
        return 0, ""

    monkeypatch.setattr(office, "com_apps", lambda: [])
    monkeypatch.setattr(office, "libreoffice_path", lambda: "C:/fake/soffice.exe")
    monkeypatch.setattr(office, "_run_child", fake_child)

    result = office.to_pdf(str(src))

    assert result["ok"] and result["via"] == "LibreOffice"
    out = Path(result["path"])
    assert out.read_bytes() == b"%PDF-1.4 fake"
    assert (home / "scratch") in out.parents
    assert not [p for p in (home / "scratch").rglob("*") if p.name.startswith(".lo-")]


def test_child_environment_drops_credential_like_names(monkeypatch):
    monkeypatch.setenv("COWORKER_TEST_API_KEY", "x")
    monkeypatch.setenv("TELEGRAM_BOT_NAME", "x")
    monkeypatch.setenv("COWORKER_PLAIN_VALUE", "kept")
    env = office._child_env()
    assert "COWORKER_TEST_API_KEY" not in env
    assert "TELEGRAM_BOT_NAME" not in env
    assert env["COWORKER_PLAIN_VALUE"] == "kept"


# ----------------------------------------------------------- PDF outputs

def test_to_pdf_output_lands_in_scratch_not_beside_the_source(tmp_path, home, monkeypatch):
    src = tmp_path / "docs" / "letter.docx"
    src.parent.mkdir()
    src.write_bytes(b"docx")
    before = listing(src.parent)

    def fake_com(source, target):
        target.write_bytes(b"%PDF-1.7 fake")
        return {"ok": True, "path": str(target), "via": "Office/Word"}

    monkeypatch.setattr(office, "com_apps", lambda: ["Word"])
    monkeypatch.setattr(office, "_pdf_via_com", fake_com)

    result = office.to_pdf(str(src))

    assert result["ok"] and result["reused"] is False
    out = Path(result["path"])
    assert (home / "scratch") in out.parents
    assert out.name == "letter.pdf"
    assert listing(src.parent) == before
    assert src.read_bytes() == b"docx"


def test_fresh_output_is_reused_and_an_edited_source_is_rendered_again(tmp_path, monkeypatch):
    src = tmp_path / "letter.docx"
    src.write_bytes(b"docx")
    calls: list[int] = []

    def fake_com(source, target):
        calls.append(1)
        target.write_bytes(b"%PDF fake")
        return {"ok": True, "path": str(target), "via": "Office/Word"}

    monkeypatch.setattr(office, "com_apps", lambda: ["Word"])
    monkeypatch.setattr(office, "_pdf_via_com", fake_com)

    first = office.to_pdf(str(src))
    second = office.to_pdf(str(src))
    assert len(calls) == 1
    assert second["reused"] is True and second["path"] == first["path"]

    src.write_bytes(b"docx, edited")
    third = office.to_pdf(str(src))
    assert len(calls) == 2
    assert third["path"] != first["path"]


def test_freshness_rejects_empty_and_older_outputs(tmp_path):
    source = tmp_path / "a.docx"
    source.write_bytes(b"x")
    target = tmp_path / "a.pdf"
    base = 1_700_000_000
    os.utime(source, (base + 100, base + 100))

    target.write_bytes(b"")
    os.utime(target, (base + 200, base + 200))
    assert office._is_fresh(target, source) is False        # empty

    target.write_bytes(b"%PDF")
    os.utime(target, (base, base))
    assert office._is_fresh(target, source) is False        # older than the source

    os.utime(target, (base + 200, base + 200))
    assert office._is_fresh(target, source) is True


def test_pdf_input_is_copied_into_scratch_and_the_source_is_unchanged(tmp_path, home):
    src = make_pdf(tmp_path / "scan.pdf", 2)
    original = src.read_bytes()

    result = office.to_pdf(str(src))

    out = Path(result["path"])
    assert result["ok"] and (home / "scratch") in out.parents
    assert out.read_bytes() == original
    assert src.read_bytes() == original


def test_pdf_pages_keeps_requested_order_and_names_from_the_parsed_list(tmp_path, home):
    src = make_pdf(tmp_path / "Report.pdf", 7)

    result = office.pdf_pages(str(src), "3,1-2,7")

    assert result["ok"] and result["pages"] == 4 and result["of"] == 7
    out = Path(result["path"])
    assert out.name == "Report_p3_1-2_7.pdf"
    assert (home / "scratch") in out.parents
    with pymupdf.open(out) as doc:
        texts = [doc[i].get_text().strip() for i in range(doc.page_count)]
    assert texts == ["page 3", "page 1", "page 2", "page 7"]


@pytest.mark.parametrize("spec", ["..\\..\\evil", "1,abc", "", "99", "1-", "../x"])
def test_bad_page_specs_are_refused_and_write_nothing(tmp_path, home, spec):
    src = make_pdf(tmp_path / "Report.pdf", 3)
    before = listing(tmp_path)

    result = office.pdf_pages(str(src), spec)

    assert result["code"] == "arg_invalid"
    assert listing(tmp_path) == before
    assert not (home / "scratch").exists()


def test_pdf_pages_reuses_a_fresh_output(tmp_path):
    src = make_pdf(tmp_path / "Report.pdf", 3)
    first = office.pdf_pages(str(src), "2")
    second = office.pdf_pages(str(src), "2")
    assert second["reused"] is True
    assert second["path"] == first["path"]


# ------------------------------------------------------------- pure helpers

@pytest.mark.parametrize("spec, total, expected", [
    ("1-3,7", 10, [0, 1, 2, 6]),
    ("3,1-2,7", 10, [2, 0, 1, 6]),
    ("2-99", 5, [1, 2, 3, 4]),
    ("7-5", 6, [4, 5]),
    ("1,1", 3, [0]),
    ("0,4", 3, []),
    ("", 3, []),
])
def test_parse_pages_is_one_based_clamped_and_ordered(spec, total, expected):
    assert office._parse_pages(spec, total) == expected


@pytest.mark.parametrize("spec", ["1-3-5", "a", "-2", "2-", "1;3"])
def test_malformed_page_specs_are_rejected_not_skipped(spec):
    assert office._parse_pages(spec, 5) is None


def test_page_token_collapses_runs_and_digests_long_lists():
    assert office._page_token([0, 1, 2, 6]) == "1-3_7"
    token = office._page_token(list(range(0, 200, 2)))
    assert token.startswith("100pages-") and len(token) < 40


@pytest.mark.parametrize("raw, expected", [
    ("1 500", 1500),
    ("1,5", 1.5),
    ("12", 12),
    ("nan", "nan"),
    ("inf", "inf"),
    ("Rent", "Rent"),
    (2.5, 2.5),
    (None, None),
    (True, True),
])
def test_coerce_keeps_numbers_numeric_and_non_finite_text(raw, expected):
    value = office._coerce(raw)
    assert value == expected
    assert type(value) is type(expected)


def test_capabilities_report_the_backends_and_abandoned_com_calls():
    caps = office.capabilities()
    assert {"read_sheets", "office_com", "libreoffice", "pdf_to_pdf", "pdf_pages", "com_abandoned"} <= caps.keys()
    assert isinstance(caps["com_abandoned"], int)
