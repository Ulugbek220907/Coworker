"""Document operations that need no mouse, no screenshot and no focus.

This is the layer Microsoft's own UFO research calls "native API preferred,
GUI fallback", and it is deliberately built first: it is deterministic,
testable, works while the screen is locked, and covers most of what an office
assistant is actually asked to do.

Two backends, chosen per operation:

  pure Python (openpyxl / python-docx / PyMuPDF)
      Always available, works on a machine with no Office installed, and is
      what gets exercised by the tests. Reading and page surgery live here.

  COM (pywin32 + a real Office install)
      Only reached for things pure Python genuinely cannot do faithfully:
      rendering a .docx to PDF with correct layout, and the legacy .doc/.xls
      formats. LibreOffice is tried next, and if neither exists the caller is
      told plainly rather than handed a bad conversion.

Anything that modifies a file takes a backup first. A remote agent editing a
contract with no undo is not something to be relaxed about.
"""
from __future__ import annotations

import logging
import os
import shutil
import subprocess
import time
from datetime import datetime
from pathlib import Path
from typing import Any

log = logging.getLogger("office")

MAX_ROWS = 40
MAX_COLS = 15
CONVERT_TIMEOUT = 180

SHEET_EXT = {".xlsx", ".xlsm", ".xltx"}
LEGACY_EXT = {".doc", ".xls", ".ppt"}
CONVERTIBLE = {".docx", ".xlsx", ".pptx", ".doc", ".xls", ".ppt", ".odt", ".ods", ".rtf", ".txt"}


# --------------------------------------------------------------- backends

def has_openpyxl() -> bool:
    import importlib.util
    return bool(importlib.util.find_spec("openpyxl"))


def com_apps() -> list[str]:
    """Which Office applications are actually registered on this machine."""
    if os.name != "nt":
        return []
    import importlib.util
    if not importlib.util.find_spec("win32com"):
        return []
    import winreg

    found = []
    for prog in ("Word.Application", "Excel.Application", "PowerPoint.Application"):
        try:
            winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, prog).Close()
            found.append(prog.split(".")[0])
        except OSError:
            continue
    return found


def libreoffice_path() -> str:
    for var in ("ProgramFiles", "ProgramFiles(x86)"):
        candidate = Path(os.environ.get(var, "")) / "LibreOffice" / "program" / "soffice.exe"
        if candidate.is_file():
            return str(candidate)
    return shutil.which("soffice") or shutil.which("libreoffice") or ""


def capabilities() -> dict:
    """Reported to the user so a missing backend is visible, not mysterious."""
    apps = com_apps()
    lo = libreoffice_path()
    return {
        "read_sheets": has_openpyxl(),
        "office_com": apps,
        "libreoffice": bool(lo),
        "pdf_to_pdf": bool(apps) or bool(lo),
        "pdf_pages": _has_fitz(),
    }


def _has_fitz() -> bool:
    import importlib.util
    return bool(importlib.util.find_spec("pymupdf") or importlib.util.find_spec("fitz"))


def _pymupdf():
    """`fitz` is the legacy alias and now warns on import."""
    try:
        import pymupdf
        return pymupdf
    except ImportError:
        import fitz
        return fitz


# ------------------------------------------------------------------- read

def sheet_list(path: str) -> dict:
    """Sheet names with their used size - cheap way to pick the right tab."""
    p = Path(path)
    if p.suffix.lower() not in SHEET_EXT:
        return {"error": f"Bu Excel fayli emas: {p.suffix}"}
    if not has_openpyxl():
        return {"error": "openpyxl o'rnatilmagan"}

    from openpyxl import load_workbook

    try:
        wb = load_workbook(p, read_only=True, data_only=True)
    except Exception as exc:
        return {"error": f"Ochib bo'lmadi: {exc}"}
    try:
        return {
            "path": str(p),
            "sheets": [
                {"name": ws.title, "rows": ws.max_row or 0, "cols": ws.max_column or 0}
                for ws in wb.worksheets
            ],
        }
    finally:
        wb.close()


def sheet_read(
    path: str,
    sheet: str | None = None,
    max_rows: int = MAX_ROWS,
    max_cols: int = MAX_COLS,
) -> dict:
    """A worksheet as a compact text table, with formulas already evaluated.

    ``data_only=True`` returns the value Excel last cached, not the formula -
    which is what a person asking "how much was the profit" actually wants.
    """
    p = Path(path)
    if p.suffix.lower() not in SHEET_EXT:
        return {"error": f"Bu Excel fayli emas: {p.suffix}"}
    if not has_openpyxl():
        return {"error": "openpyxl o'rnatilmagan"}

    from openpyxl import load_workbook

    try:
        wb = load_workbook(p, read_only=True, data_only=True)
    except Exception as exc:
        return {"error": f"Ochib bo'lmadi: {exc}"}

    try:
        if sheet:
            match = _find_sheet(wb, sheet)
            if match is None:
                return {"error": f"«{sheet}» varag'i yo'q. Bor: {', '.join(wb.sheetnames)}"}
            ws = match
        else:
            ws = wb.worksheets[0]

        rows: list[list[str]] = []
        truncated_cols = False
        for row in ws.iter_rows(max_row=max_rows + 1, max_col=max_cols + 1, values_only=True):
            if len(rows) >= max_rows:
                break
            if row and len(row) > max_cols:
                truncated_cols = True
                row = row[:max_cols]
            cells = ["" if v is None else _fmt(v) for v in (row or ())]
            if any(cells):
                rows.append(cells)

        return {
            "path": str(p),
            "sheet": ws.title,
            "all_sheets": wb.sheetnames,
            "table": _as_text(rows),
            "shown_rows": len(rows),
            "total_rows": ws.max_row or 0,
            "truncated": (ws.max_row or 0) > max_rows or truncated_cols,
        }
    except Exception as exc:
        return {"error": f"O'qishda xato: {exc}"}
    finally:
        wb.close()


def _find_sheet(wb, wanted: str):
    """Match a sheet name across scripts - the tab may be Cyrillic."""
    from .textutil import normalize

    target = normalize(wanted)
    for ws in wb.worksheets:
        if ws.title == wanted or normalize(ws.title) == target:
            return ws
    for ws in wb.worksheets:                      # then a loose contains-match
        if target and target in normalize(ws.title):
            return ws
    return None


def _fmt(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int) and abs(value) >= 10000:
        return f"{value:,}".replace(",", " ")     # 1 240 000 000 reads far better
    return str(value)


def _as_text(rows: list[list[str]]) -> str:
    """Pipe-separated, column-aligned - compact for a model and for a phone."""
    if not rows:
        return "(bo'sh)"
    # iter_rows pads every row out to max_col, so drop the trailing empty
    # columns before measuring - otherwise a two-column sheet renders with a
    # dozen empty pipes and burns tokens for nothing.
    width = 0
    for r in rows:
        last = max((i + 1 for i, c in enumerate(r) if c), default=0)
        width = max(width, last)
    if not width:
        return "(bo'sh)"
    rows = [(r + [""] * width)[:width] for r in rows]
    widths = [min(24, max(len(r[i]) for r in rows)) for i in range(width)]
    return "\n".join(
        " | ".join(cell[:24].ljust(w) for cell, w in zip(row, widths)).rstrip()
        for row in rows
    )


# -------------------------------------------------------- produce new files

def to_pdf(path: str, out_dir: str | None = None) -> dict:
    """Render a document to PDF. Never modifies the source."""
    src = Path(path)
    if not src.is_file():
        return {"error": f"Fayl yo'q: {path}"}
    if src.suffix.lower() == ".pdf":
        return {"ok": True, "path": str(src), "note": "allaqachon PDF"}
    if src.suffix.lower() not in CONVERTIBLE:
        return {"error": f"Bu format PDF'ga o'girilmaydi: {src.suffix}"}

    out = Path(out_dir) if out_dir else _work_dir()
    out.mkdir(parents=True, exist_ok=True)
    target = out / (src.stem + ".pdf")

    for backend in (_pdf_via_com, _pdf_via_libreoffice):
        result = backend(src, target)
        if result.get("ok"):
            return result
        if result.get("error"):
            log.info("%s: %s", backend.__name__, result["error"])

    return {
        "error": (
            "PDF'ga o'girish uchun Microsoft Office yoki LibreOffice kerak. "
            "Bu kompyuterda ikkalasi ham yo'q."
        )
    }


def _pdf_via_com(src: Path, target: Path) -> dict:
    """Word/Excel/PowerPoint via COM. Highest fidelity when Office is present."""
    apps = com_apps()
    ext = src.suffix.lower()
    kind = (
        "Word" if ext in {".docx", ".doc", ".rtf", ".odt", ".txt"}
        else "Excel" if ext in {".xlsx", ".xls", ".ods"}
        else "PowerPoint" if ext in {".pptx", ".ppt"}
        else ""
    )
    if not kind or kind not in apps:
        return {"error": f"{kind or ext} uchun COM yo'q"}

    import pythoncom
    import win32com.client

    # Tool calls run on a worker thread; COM demands per-thread initialisation.
    pythoncom.CoInitialize()
    app = None
    doc = None
    try:
        app = win32com.client.DispatchEx(f"{kind}.Application")
        try:
            app.Visible = False
        except Exception:
            pass                                   # PowerPoint refuses to hide
        try:
            app.DisplayAlerts = False
        except Exception:
            pass

        full_src, full_target = str(src.resolve()), str(target.resolve())
        if kind == "Word":
            doc = app.Documents.Open(full_src, ReadOnly=True)
            doc.ExportAsFixedFormat(full_target, 17)          # 17 = wdExportFormatPDF
        elif kind == "Excel":
            doc = app.Workbooks.Open(full_src, ReadOnly=True)
            doc.ExportAsFixedFormat(0, full_target)           # 0 = xlTypePDF
        else:
            doc = app.Presentations.Open(full_src, ReadOnly=True, WithWindow=False)
            doc.SaveAs(full_target, 32)                       # 32 = ppSaveAsPDF

        if not target.is_file():
            return {"error": "Office xatosiz tugadi, lekin fayl yaratilmadi"}
        return {"ok": True, "path": str(target), "via": f"Office/{kind}"}

    except Exception as exc:
        return {"error": f"COM: {exc}"}
    finally:
        # Leaking a hidden EXCEL.EXE on every call is the classic failure here.
        for obj, method in ((doc, "Close"), (app, "Quit")):
            if obj is not None:
                try:
                    getattr(obj, method)()
                except Exception:
                    pass
        pythoncom.CoUninitialize()


def _pdf_via_libreoffice(src: Path, target: Path) -> dict:
    exe = libreoffice_path()
    if not exe:
        return {"error": "LibreOffice yo'q"}
    try:
        proc = subprocess.run(
            [exe, "--headless", "--norestore", "--convert-to", "pdf",
             "--outdir", str(target.parent), str(src.resolve())],
            capture_output=True, timeout=CONVERT_TIMEOUT,
        )
    except subprocess.TimeoutExpired:
        return {"error": "LibreOffice javob bermadi"}
    if not target.is_file():
        return {"error": (proc.stderr or b"").decode(errors="replace")[:200] or "natija yo'q"}
    return {"ok": True, "path": str(target), "via": "LibreOffice"}


def pdf_pages(path: str, pages: str, out_dir: str | None = None) -> dict:
    """Pull a page range out of a PDF - "send me just page 3" on a phone."""
    src = Path(path)
    if not src.is_file() or src.suffix.lower() != ".pdf":
        return {"error": "PDF fayl kerak"}
    if not _has_fitz():
        return {"error": "PyMuPDF o'rnatilmagan"}

    fitz = _pymupdf()

    try:
        with fitz.open(src) as doc:
            total = doc.page_count
            wanted = _parse_pages(pages, total)
            if not wanted:
                return {"error": f"Sahifa raqami noto'g'ri. Faylda {total} sahifa bor."}
            out = Path(out_dir) if out_dir else _work_dir()
            out.mkdir(parents=True, exist_ok=True)
            target = out / f"{src.stem}_{pages.replace(',', '_').replace(' ', '')}.pdf"
            with fitz.open() as new:
                for index in wanted:
                    new.insert_pdf(doc, from_page=index, to_page=index)
                new.save(str(target))
    except Exception as exc:
        return {"error": f"PDF xatosi: {exc}"}

    return {"ok": True, "path": str(target), "pages": len(wanted), "of": total}


def _parse_pages(spec: str, total: int) -> list[int]:
    """"1-3,7" -> [0,1,2,6]. One-based in, zero-based out, clamped."""
    out: list[int] = []
    for part in str(spec).replace(" ", "").split(","):
        if not part:
            continue
        try:
            if "-" in part:
                lo, hi = part.split("-", 1)
                start, end = int(lo), int(hi)
            else:
                start = end = int(part)
        except ValueError:
            continue
        for n in range(min(start, end), max(start, end) + 1):
            if 1 <= n <= total and (n - 1) not in out:
                out.append(n - 1)
    return out


# ---------------------------------------------------------------- modify

def sheet_write(path: str, changes: dict, sheet: str | None = None) -> dict:
    """Set cells, e.g. {"B5": 500000}. Backs the file up first.

    openpyxl round-trips a workbook by rewriting it, which can drop charts,
    images and some formatting. The backup is not a nicety - it is the undo.
    """
    p = Path(path)
    if p.suffix.lower() not in SHEET_EXT:
        return {"error": "Faqat .xlsx fayllarini tahrirlay olaman"}
    if not has_openpyxl():
        return {"error": "openpyxl o'rnatilmagan"}
    if not changes:
        return {"error": "O'zgarishlar ko'rsatilmagan"}

    backup = _backup(p)
    if backup.get("error"):
        return backup

    from openpyxl import load_workbook

    try:
        wb = load_workbook(p)                      # keep formulas, so not data_only
        ws = _find_sheet(wb, sheet) if sheet else wb.worksheets[0]
        if ws is None:
            return {"error": f"«{sheet}» varag'i yo'q"}

        applied = {}
        for ref, value in changes.items():
            ws[str(ref).upper()] = _coerce(value)
            applied[str(ref).upper()] = value
        wb.save(p)
        wb.close()
    except Exception as exc:
        return {"error": f"Yozishda xato: {exc}", "backup": backup["backup"]}

    return {"ok": True, "path": str(p), "sheet": ws.title,
            "applied": applied, "backup": backup["backup"]}


def _coerce(value: Any) -> Any:
    """A model sends JSON; a spreadsheet wants numbers to stay numbers."""
    if isinstance(value, (int, float, bool)) or value is None:
        return value
    text = str(value).strip()
    cleaned = text.replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        return int(cleaned)
    except ValueError:
        pass
    try:
        return float(cleaned)
    except ValueError:
        return text


def _backup(p: Path) -> dict:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    target = p.with_name(f"{p.stem}.backup-{stamp}{p.suffix}")
    try:
        shutil.copy2(p, target)
    except OSError as exc:
        return {"error": f"Zaxira nusxa olinmadi, o'zgartirmadim: {exc}"}
    return {"backup": str(target)}


def _work_dir() -> Path:
    """Generated files live beside the config, never next to the user's data."""
    from .config import config_dir

    d = config_dir() / "generated"
    d.mkdir(parents=True, exist_ok=True)
    _prune(d)
    return d


def _prune(d: Path, keep_hours: int = 48) -> None:
    cutoff = time.time() - keep_hours * 3600
    try:
        for f in d.iterdir():
            if f.is_file() and f.stat().st_mtime < cutoff:
                f.unlink(missing_ok=True)
    except OSError:
        pass
